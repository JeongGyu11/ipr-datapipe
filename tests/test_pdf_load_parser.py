"""PDF 파서의 원문 보존·헤더푸터·페이지 품질 테스트."""

from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

from app.domain.pdf_load import PageStatus
import app.infrastructure.pdf.prosure.parser as parser_module
from app.infrastructure.pdf.prosure.header_footer import normalize_candidate
from app.infrastructure.pdf.prosure.parser import (
    ProSurePdfParser,
    _extract_text_objects,
    _join_page_objects,
    _normalise_cell,
    _page_quality,
    _sanitize_table_rows,
)


def _write_text_pdf(
    path: Path,
    *,
    page_count: int = 1,
    different_first_header: bool = False,
) -> None:
    document = pymupdf.open()
    for page_no in range(1, page_count + 1):
        page = document.new_page()
        header = "Cover Title" if different_first_header and page_no == 1 else "Insurance Product Header"
        page.insert_text((72, 20), header)
        page.insert_text((72, 100), f"Insurance policy sample body page {page_no}.")
        page.insert_text((72, 820), f"Page {page_no} / {page_count}")
    document.save(path)
    document.close()


def _write_image_pdf(path: Path, *, include_text_page: bool = False) -> None:
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10), False)
    pixmap.clear_with(255)
    pixel = pixmap.tobytes("png")
    document = pymupdf.open()
    if include_text_page:
        page = document.new_page()
        page.insert_text((72, 100), "Insurance policy usable text page 1.")
    page = document.new_page()
    page.insert_image(page.rect, stream=pixel)
    document.save(path)
    document.close()


def _write_image_pdf_with_repeating_header(path: Path, *, page_count: int = 5) -> None:
    """전체 페이지 이미지 위에 반복 헤더를 얹은 OCR 품질 fixture."""

    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10), False)
    pixmap.clear_with(255)
    pixel = pixmap.tobytes("png")
    document = pymupdf.open()
    for _ in range(page_count):
        page = document.new_page()
        page.insert_image(page.rect, stream=pixel)
        page.insert_text((72, 20), "Repeated OCR Header")
    document.save(path)
    document.close()


def test_header_footer_off_preserves_text_and_does_not_create_crop(tmp_path: Path) -> None:
    source = tmp_path / "sample.pdf"
    _write_text_pdf(source)
    item_dir = tmp_path / "item"

    output = ProSurePdfParser(header_footer_enabled=False).parse(source, item_dir)

    page = output.pages[0]
    assert "Insurance Product Header" in page.text_raw
    assert "Insurance Product Header" in page.text_content
    assert page.content_markdown
    assert page.text_raw
    assert output.header_footer_analysis["enabled"] is False
    assert output.header_footer_analysis["applied"] is False
    assert output.cropped_pdf_path is None
    assert not (item_dir / "90_debug" / "cropped.pdf").exists()


def test_header_footer_on_removes_only_repeated_lines_and_keeps_raw(tmp_path: Path) -> None:
    source = tmp_path / "repeated.pdf"
    _write_text_pdf(source, page_count=5)
    item_dir = tmp_path / "item"

    output = ProSurePdfParser(header_footer_enabled=True).parse(source, item_dir)

    assert output.header_footer_analysis["applied"] is True
    assert all("Insurance Product Header" in page.text_raw for page in output.pages)
    assert all("Insurance Product Header" not in page.text_content for page in output.pages)
    assert all("Insurance policy sample body" in page.text_content for page in output.pages)
    assert output.cropped_pdf_path == Path("90_debug/cropped.pdf")
    assert (item_dir / output.cropped_pdf_path).exists()
    assert output.header_footer_analysis["removed_lines"]


def test_header_footer_on_does_not_remove_lines_from_short_document(tmp_path: Path) -> None:
    source = tmp_path / "short.pdf"
    _write_text_pdf(source, page_count=2)

    output = ProSurePdfParser(header_footer_enabled=True).parse(source, tmp_path / "item")

    assert output.header_footer_analysis["applied"] is False
    assert all("Insurance Product Header" in page.text_content for page in output.pages)
    assert output.cropped_pdf_path is None


def test_header_footer_keeps_non_repeated_first_page_line(tmp_path: Path) -> None:
    source = tmp_path / "different-first.pdf"
    _write_text_pdf(source, page_count=5, different_first_header=True)

    output = ProSurePdfParser(header_footer_enabled=True).parse(source, tmp_path / "item")

    assert "Cover Title" in output.pages[0].text_content
    assert all("Insurance Product Header" not in page.text_content for page in output.pages[1:])
    with pymupdf.open(tmp_path / "item" / output.cropped_pdf_path) as cropped:
        assert cropped[0].rect.height > cropped[1].rect.height


def test_page_markdown_has_one_blank_line_before_and_after_table() -> None:
    page_view = _join_page_objects(
        [
            {"y": 10.0, "content": "표 앞 본문"},
            {"y": 20.0, "content": "| 항목 | 값 |\n| --- | --- |", "kind": "TABLE"},
            {"y": 30.0, "content": "표 뒤 본문"},
            {"y": 40.0, "content": "다음 본문"},
        ]
    )

    assert page_view == (
        "표 앞 본문\n\n"
        "| 항목 | 값 |\n"
        "| --- | --- |\n\n"
        "표 뒤 본문\n"
        "다음 본문"
    )


def test_table_sanitization_removes_only_excel_xml_illegal_controls() -> None:
    rows, removed_count, affected_cells = _sanitize_table_rows(
        [["a\x00b\t c", "line\nfeed\rcarriage"], ["\x0b\x0c\x1f", "ok"]]
    )

    assert rows == [["ab\t c", "line\nfeed\rcarriage"], ["", "ok"]]
    assert removed_count == 4
    assert affected_cells == [
        {"row": 1, "column": 1, "removed_count": 1},
        {"row": 2, "column": 1, "removed_count": 3},
    ]


def test_cell_normalization_preserves_allowed_line_controls() -> None:
    assert _normalise_cell(" \tvalue\n\r ") == "\tvalue\n\r"


def test_parser_sanitizes_table_artifacts_and_records_warning_without_touching_raw(
    tmp_path: Path, monkeypatch
) -> None:
    source = tmp_path / "sample.pdf"
    source.write_bytes(b"not a real pdf")
    raw_text = "raw\x00pdf text"

    class FakeTable:
        bbox = (0, 0, 100, 100)

        @staticmethod
        def extract():
            return [["header", "bad\x01value\tline\nfeed\rcarriage"], ["ok", "clean"]]

    class FakePage:
        width = 200
        height = 300
        images: list[dict[str, object]] = []

        @staticmethod
        def extract_text():
            return raw_text

        @staticmethod
        def find_tables():
            return [FakeTable()]

        @staticmethod
        def extract_words():
            return [{"text": "body", "x0": 150, "x1": 180, "top": 150, "bottom": 160}]

    class FakePdf:
        pages = [FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

    class FakePdfPlumber:
        @staticmethod
        def open(path):
            return FakePdf()

    monkeypatch.setattr(parser_module, "_require_pdfplumber", lambda: FakePdfPlumber)
    monkeypatch.setattr(
        parser_module,
        "extract_embedded_images",
        lambda source_pdf, output_dir: ([], {}, []),
    )

    output = ProSurePdfParser().parse(source, tmp_path / "item")

    warning = next(
        warning
        for warning in output.warnings
        if warning["code"] == "TABLE_ILLEGAL_CONTROL_CHARACTERS_REMOVED"
    )
    assert warning == {
        "code": "TABLE_ILLEGAL_CONTROL_CHARACTERS_REMOVED",
        "page_no": 1,
        "table_id": "page_0001_table_001",
        "table_index": 1,
        "candidate_index": 1,
        "removed_count": 1,
        "affected_cells": [{"row": 1, "column": 2, "removed_count": 1}],
    }
    assert output.pages[0].text_raw == raw_text
    assert output.raw_text == raw_text
    xlsx_path = tmp_path / "item" / "02_tables" / "page_0001_table_001.xlsx"
    workbook = load_workbook(xlsx_path, read_only=True)
    try:
        assert list(workbook.active.values) == [
            ("header", "badvalue\tline\nfeed\rcarriage"),
            ("ok", "clean"),
        ]
    finally:
        workbook.close()
    markdown = (tmp_path / "item" / "02_tables" / "page_0001_table_001.md").read_text(
        encoding="utf-8"
    )
    assert markdown == (
        "| header | badvalue\tline feed carriage |\n"
        "| --- | --- |\n"
        "| ok | clean |"
    )


@pytest.mark.parametrize("failure_kind", ["text", "detection", "extraction", "render"])
def test_parser_warning_records_have_stable_codes_and_separate_details(
    tmp_path: Path, monkeypatch, failure_kind: str
) -> None:
    source = tmp_path / f"{failure_kind}.pdf"
    source.write_bytes(b"not a real pdf")

    class FakeTable:
        bbox = (0, 0, 100, 100)

        @staticmethod
        def extract():
            if failure_kind == "extraction":
                raise RuntimeError("extract: boom")
            return [["header", "value"], ["보험료", "100만원"]]

    class FakePage:
        width = 200
        height = 300
        images: list[dict[str, object]] = []

        @staticmethod
        def extract_text():
            if failure_kind == "text":
                raise RuntimeError("text: boom")
            return "body"

        @staticmethod
        def find_tables():
            if failure_kind == "detection":
                raise RuntimeError("detect: boom")
            return [FakeTable()] if failure_kind in {"extraction", "render"} else []

        @staticmethod
        def extract_words():
            return [{"text": "body", "x0": 150, "x1": 180, "top": 150, "bottom": 160}]

    class FakePdf:
        pages = [FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

    class FakePdfPlumber:
        @staticmethod
        def open(path):
            return FakePdf()

    class FailingRenderer:
        def __init__(self, source_pdf):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        @staticmethod
        def render(page_no, bbox, output_path):
            raise RuntimeError("render: boom")

    monkeypatch.setattr(parser_module, "_require_pdfplumber", lambda: FakePdfPlumber)
    monkeypatch.setattr(
        parser_module,
        "extract_embedded_images",
        lambda source_pdf, output_dir: ([], {}, []),
    )
    if failure_kind == "render":
        monkeypatch.setattr(parser_module, "TableRegionRenderer", FailingRenderer)

    output = ProSurePdfParser().parse(source, tmp_path / "item")

    expected_code = {
        "text": "TEXT_EXTRACTION_FAILED",
        "detection": "TABLE_DETECTION_FAILED",
        "extraction": "TABLE_EXTRACTION_FAILED",
        "render": "TABLE_RENDER_FAILED",
    }[failure_kind]
    record = next(warning for warning in output.warnings if warning["code"] == expected_code)
    assert record["page_no"] == 1
    expected_message_prefix = {
        "text": "text",
        "detection": "detect",
        "extraction": "extract",
        "render": "render",
    }[failure_kind]
    assert record["message"] == f"{expected_message_prefix}: boom"
    assert expected_code in output.pages[0].warnings
    assert ":" not in str(record["code"])
    if failure_kind in {"extraction", "render"}:
        assert record["candidate_index"] == 1
    if failure_kind == "render":
        assert record["table_index"] == 1


def test_parser_rejects_single_cell_layout_and_preserves_it_as_text(
    tmp_path: Path, monkeypatch
) -> None:
    source = tmp_path / "sample.pdf"
    source.write_bytes(b"not a real pdf")

    class FakeTable:
        def __init__(self, bbox, rows):
            self.bbox = bbox
            self.rows = rows

        def extract(self):
            return self.rows

    class FakePage:
        width = 200
        height = 300
        images: list[dict[str, object]] = []

        @staticmethod
        def extract_text():
            return "상 품 요 약 서\n항목 값\n보험료 100만원\n본문"

        @staticmethod
        def find_tables():
            return [
                FakeTable((10, 10, 190, 50), [["상 품 요 약 서"]]),
                FakeTable(
                    (10, 100, 190, 180),
                    [["항목", "값"], ["보험료 20년", "100만원 50,000,000"]],
                ),
                FakeTable((80, 140, 180, 170), [["20년", "50,000,000"]]),
            ]

        @staticmethod
        def extract_words():
            return [
                {"text": "상", "x0": 20, "x1": 30, "top": 20, "bottom": 30},
                {"text": "품", "x0": 40, "x1": 50, "top": 20, "bottom": 30},
                {"text": "요", "x0": 60, "x1": 70, "top": 20, "bottom": 30},
                {"text": "약", "x0": 80, "x1": 90, "top": 20, "bottom": 30},
                {"text": "서", "x0": 100, "x1": 110, "top": 20, "bottom": 30},
                {"text": "항목", "x0": 20, "x1": 50, "top": 120, "bottom": 130},
                {"text": "값", "x0": 110, "x1": 130, "top": 120, "bottom": 130},
                {"text": "보험료", "x0": 20, "x1": 60, "top": 150, "bottom": 160},
                {"text": "100만원", "x0": 110, "x1": 160, "top": 150, "bottom": 160},
                {"text": "본문", "x0": 20, "x1": 50, "top": 220, "bottom": 230},
            ]

    class FakePdf:
        pages = [FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

    class FakePdfPlumber:
        @staticmethod
        def open(path):
            return FakePdf()

    class FakeRenderer:
        def __init__(self, source_pdf):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        @staticmethod
        def render(page_no, bbox, output_path):
            output_path.write_bytes(b"png")
            return output_path

    monkeypatch.setattr(parser_module, "_require_pdfplumber", lambda: FakePdfPlumber)
    monkeypatch.setattr(parser_module, "TableRegionRenderer", FakeRenderer)
    monkeypatch.setattr(
        parser_module,
        "extract_embedded_images",
        lambda source_pdf, output_dir: ([], {}, []),
    )

    output = ProSurePdfParser().parse(source, tmp_path / "item")

    assert output.detected_table_count == 3
    assert output.rejected_table_count == 1
    assert len(output.tables) == 2
    assert [table.table_id for table in output.tables] == [
        "page_0001_table_001",
        "page_0001_table_002",
    ]
    assert output.fragment_table_count == 1
    assert output.pages[0].table_refs == [
        "page_0001_table_001",
        "page_0001_table_002",
    ]
    assert "상 품 요 약 서" in output.pages[0].text_content
    assert "상 품 요 약 서" in output.pages[0].content_markdown
    assert "본문" in output.pages[0].text_content
    assert "100만원" not in output.pages[0].text_content
    assert (tmp_path / "item/02_tables/page_0001_table_002.md").exists()
    assert not (tmp_path / "item/02_tables/page_0001_table_003.md").exists()
    warning = next(
        warning
        for warning in output.warnings
        if warning["code"] == "TABLE_CANDIDATE_REJECTED"
    )
    assert warning["reason"] == "SINGLE_CELL_LAYOUT"
    assert warning["candidate_index"] == 1
    assert warning["metrics"]["nonempty_cell_count"] == 1


def test_text_objects_exclude_table_and_only_explicit_removed_line() -> None:
    class FakePage:
        @staticmethod
        def extract_words():
            return [
                {"text": "HEADER", "x0": 10, "x1": 60, "top": 4, "bottom": 10},
                {"text": "본문앞", "x0": 10, "x1": 60, "top": 100, "bottom": 110},
                {"text": "표중복", "x0": 100, "x1": 150, "top": 300, "bottom": 310},
                {"text": "본문뒤", "x0": 10, "x1": 60, "top": 500, "bottom": 510},
                {"text": "FOOTER", "x0": 10, "x1": 60, "top": 960, "bottom": 970},
            ]

    objects = _extract_text_objects(
        FakePage(),
        [(90, 290, 160, 320)],
        removed_line_bboxes=[(10, 4, 60, 10)],
        fallback_text="",
    )

    assert objects == [
        {"y": 100.0, "x": 10.0, "content": "본문앞", "kind": "TEXT"},
        {"y": 500.0, "x": 10.0, "content": "본문뒤", "kind": "TEXT"},
        {"y": 960.0, "x": 10.0, "content": "FOOTER", "kind": "TEXT"},
    ]


def test_text_content_keeps_pipe_and_only_markdown_escapes_it(tmp_path: Path) -> None:
    source = tmp_path / "pipe-text.pdf"
    document = pymupdf.open()
    try:
        page = document.new_page()
        page.insert_text((72, 72), "Benefit A|B")
        document.save(source)
    finally:
        document.close()

    output = ProSurePdfParser().parse(source, tmp_path / "item")

    assert "Benefit A|B" in output.pages[0].text_raw
    assert "Benefit A|B" in output.pages[0].text_content
    assert r"Benefit A\|B" in output.pages[0].content_markdown


def test_page_quality_distinguishes_text_ocr_and_empty() -> None:
    class FakePage:
        width = 100
        height = 100

        def __init__(self, images):
            self.images = images

        @staticmethod
        def extract_words():
            return []

    scanned = FakePage([{"x0": 0, "x1": 100, "top": 0, "bottom": 100}])
    empty = FakePage([])
    text = FakePage([])

    assert _page_quality(scanned, "잡음")[4] is PageStatus.OCR_REQUIRED
    assert _page_quality(empty, "")[4] is PageStatus.EMPTY
    assert _page_quality(empty, "", extraction_failed=True)[4] is PageStatus.OCR_REQUIRED
    assert _page_quality(text, "보험상품 본문 text 1234567890")[4] is PageStatus.TEXT


def test_real_image_only_and_mixed_pdfs_get_page_level_ocr_status(tmp_path: Path) -> None:
    scanned = tmp_path / "scanned.pdf"
    mixed = tmp_path / "mixed.pdf"
    _write_image_pdf(scanned)
    _write_image_pdf(mixed, include_text_page=True)

    scanned_output = ProSurePdfParser().parse(scanned, tmp_path / "scanned-item")
    mixed_output = ProSurePdfParser().parse(mixed, tmp_path / "mixed-item")

    assert [page.page_status for page in scanned_output.pages] == [PageStatus.OCR_REQUIRED]
    assert [page.page_status for page in mixed_output.pages] == [
        PageStatus.TEXT,
        PageStatus.OCR_REQUIRED,
    ]


def test_repeating_header_is_quality_excluded_when_header_footer_is_off_or_on(
    tmp_path: Path,
) -> None:
    source = tmp_path / "image-with-header.pdf"
    _write_image_pdf_with_repeating_header(source)

    off = ProSurePdfParser(header_footer_enabled=False).parse(source, tmp_path / "off")
    on = ProSurePdfParser(header_footer_enabled=True).parse(source, tmp_path / "on")

    assert all(page.page_status is PageStatus.OCR_REQUIRED for page in off.pages)
    assert all(page.page_status is PageStatus.OCR_REQUIRED for page in on.pages)
    assert "Repeated OCR Header" in off.pages[0].text_raw
    assert "Repeated OCR Header" in off.pages[0].text_content
    assert "Repeated OCR Header" in on.pages[0].text_raw
    assert "Repeated OCR Header" not in on.pages[0].text_content
    assert off.header_footer_analysis["removed_lines"]
    assert on.header_footer_analysis["applied"] is True


def test_page_quality_uses_ocr_threshold_at_exactly_ten_characters() -> None:
    class ImagePage:
        width = 100
        height = 100
        images = [{"x0": 0, "x1": 100, "top": 0, "bottom": 100}]

        @staticmethod
        def extract_words():
            return []

    page = ImagePage()
    assert _page_quality(page, "1234567890")[4] is PageStatus.OCR_REQUIRED
    assert _page_quality(page, "12345678901")[4] is PageStatus.TEXT


def test_page_quality_accepts_normal_body_and_table_only_as_usable_data() -> None:
    class TextPage:
        width = 100
        height = 100
        images: list[dict[str, object]] = []

        @staticmethod
        def extract_words():
            return []

    page = TextPage()
    assert _page_quality(page, "정상적인 본문 데이터")[4] is PageStatus.TEXT
    assert _page_quality(page, "", table_present=True)[4] is PageStatus.TEXT


def test_header_candidate_normalization_ignores_page_numbers() -> None:
    assert normalize_candidate("상품 약관 1 / 10") == normalize_candidate("상품   약관 2/10")


def test_header_candidate_normalization_preserves_dates_and_amounts() -> None:
    assert normalize_candidate("기준일 2026-09-01 보험료 100만원") != normalize_candidate(
        "기준일 2025-09-01 보험료 200만원"
    )
